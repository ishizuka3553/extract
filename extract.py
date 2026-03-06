"""Word 表から日付を抽出して CSV に書き出す簡易スクリプト。"""

from typing import List, Tuple
import csv
import datetime
import logging
import time

import docx

# ダミー日付を現在日付の1年後に設定
DUMMY_DATE = datetime.datetime.now() + datetime.timedelta(days=365)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def parse_docx_tables(docx_path: str) -> List[Tuple[str, datetime.datetime]]:
    """指定した docx ファイルのテーブルから区域番号と日付を抽出して返す。

    Args:
        docx_path: 読み込む Word ファイルのパス。
    Returns:
        (区域番号, 日付) のタプルのリスト。
    """
    doc = docx.Document(docx_path)
    results: List[Tuple[str, datetime.datetime]] = []
    page_count = 1
    for tbl in doc.tables:
        if len(tbl.rows) < 2:
            continue
        logger.info("%dページ目を読み取り中", page_count)
        page_count += 1
        for row in range(2, len(tbl.rows)):
            if row % 2 == 0:
                continue
            if tbl.cell(row, 0).text.strip() == "":
                continue
            area_no = tbl.cell(row, 0).text.strip()
            found_date = None
            for col in range(1, len(tbl.columns)):
                value = tbl.cell(row, col).text.strip()
                if value != "":
                    try:
                        found_date = datetime.datetime.strptime(value, "%y/%m/%d")
                    except ValueError:
                        logger.error(
                            "エラーが発生しました。区域番号 %s の日付 %s が間違っています。",
                            area_no,
                            value,
                        )
                        found_date = None
                elif col % 2 == 1 and tbl.cell(row, col - 1).text.strip() != "":
                    found_date = DUMMY_DATE
            if found_date is not None:
                results.append((area_no, found_date))
    return results


def write_results_to_csv(
    results: List[Tuple[str, datetime.datetime]], output_path: str
) -> None:
    """抽出結果を日付でソートして CSV に書き出す。

    未返却（ダミー日付）の場合は '未返却' として出力する。
    """
    sorted_results = sorted(results, key=lambda x: x[1])
    dummy_date_str = DUMMY_DATE.strftime("%y/%m/%d")
    with open(output_path, "w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj, lineterminator="\n")
        for area_no, dt in sorted_results:
            date_str = dt.strftime("%y/%m/%d")
            if date_str == dummy_date_str:
                writer.writerow([area_no, "未返却"])
            else:
                writer.writerow([area_no, date_str])


def main():
    """
    メイン関数。
    """
    start = time.time()
    results = parse_docx_tables("北茨城・高萩区域2025.docx")
    write_results_to_csv(results, "./抽出結果.txt")
    end = time.time()
    logger.info("実行時間(秒): %.3f", end - start)


if __name__ == "__main__":
    main()
